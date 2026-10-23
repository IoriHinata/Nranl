from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as I, Pt as P
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root=Path("book-project"); out=root/"output"; out.mkdir(parents=True,exist_ok=True)
srcs=sorted(p for p in root.glob("*.md") if p.name not in ("README.md","INDEX.md"))
def plain(s):
    return s.replace("**","").replace("*","").replace("[","").replace("]","")
def build_doc(src,dest):
    d=Document(); d.sections[0].top_margin=Inches(.7); d.sections[0].bottom_margin=Inches(.7)
    d.styles["Normal"].font.name="Aptos"; d.styles["Normal"].font.size=Pt(10.5)
    lines=src.read_text(encoding="utf-8").splitlines()
    title=next((x.lstrip("# ").strip() for x in lines if x.startswith("# ")),src.stem)
    p=d.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(plain(title))
    d.add_paragraph("Город, который помнит нас · рабочие материалы")
    i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s: i+=1; continue
        if s.startswith("# "): d.add_heading(plain(s[2:]),1)
        elif s.startswith("## "): d.add_heading(plain(s[3:]),2)
        elif s.startswith("### "): d.add_heading(plain(s[4:]),3)
        elif s.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                row=[plain(x.strip()) for x in lines[i].strip().strip("|").split("|")]
                if not all(set(x.replace(":","")) <= set("- ") for x in row): rows.append(row)
                i+=1
            if rows:
                n=max(len(x) for x in rows); t=d.add_table(rows=1,cols=n); t.style="Light Shading Accent 1"
                for j in range(n): t.rows[0].cells[j].text=rows[0][j] if j<len(rows[0]) else ""
                for row in rows[1:]:
                    cells=t.add_row().cells
                    for j in range(n): cells[j].text=row[j] if j<len(row) else ""
            continue
        elif s.startswith("- ") or s.startswith("* "): d.add_paragraph(plain(s[2:]),style="List Bullet")
        elif s=="---": d.add_paragraph("— "*20)
        elif s.startswith(">"): d.add_paragraph(plain(s.lstrip("> ")),style="Quote")
        elif s.startswith(chr(96)*3): pass
        else: d.add_paragraph(plain(s))
        i+=1
    d.sections[0].footer.paragraphs[0].text="Рабочий проект · перепроверяйте актуальность источников"
    d.save(dest)
for src in srcs: build_doc(src,out/(src.stem+".docx"))
master=Document()
p=master.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Город, который помнит нас")
master.add_paragraph("Комплексный проект: исследование рынка, творческая разработка, продвижение и публикация.")
for src in srcs:
    master.add_page_break()
    tmp=out/(".temp_"+src.stem+".docx"); build_doc(src,tmp); part=Document(tmp)
    for el in list(part.element.body):
        if el.tag.endswith("sectPr"): continue
        master.element.body.insert(len(master.element.body)-1,el)
    tmp.unlink(missing_ok=True)
master.save(out/"MASTER_PROJECT_ALL_SECTIONS.docx")

prs=Presentation(); prs.slide_width=I(13.333); prs.slide_height=I(7.5)
slides=[
("Город, который помнит нас",["Комплексный проект художественного романа","Городское фэнтези / romantasy / мистический триллер","Автор свободен менять творческую основу"]),
("Логлайн",["Город стирает людей из памяти, но Мира продолжает помнить исчезнувших.","Расследование приводит её к Элиану, связанному с системой забвения.","Кто вправе избавлять общество от боли ценой чужого существования?"]),
("Обещание читателю",["Тайна с последствиями","Атмосферный город и бытовые детали-улики","Эмоциональные отношения и моральный выбор","Самостоятельная эмоциональная дуга"]),
("Тема",["Если память причиняет боль, имеет ли кто-то право её отнять?","Забвение может выглядеть как милосердие и работать как насилие.","Память связана с личностью, утратой и ответственностью."]),
("Мир и персонажи",["Мира помнит забытых; Элиан связан с механизмом забвения.","У магии должны быть ограничения, цена и следы.","У героев разные цели, ошибочные убеждения и точки невозврата."]),
("Рынок",["Художественная литература остаётся значимым сегментом.","Социальные рекомендации помогают открывать книги, но не гарантируют продажи.","Рынки, форматы и денежные показатели нельзя смешивать."]),
("Проверка концепции",["Сравнить 15–25 книг.","Проверить 3 логлайна и 2–3 визуальных направления.","Дать фрагмент целевым читателям и задать одинаковые вопросы."]),
("Создание рукописи",["Авторское ядро → черновик → проверка причин и последствий.","Бета-читатели → редактура → корректура.","Каждая сцена меняет сюжет, героя, отношения или тему."]),
("Продвижение",["Один основной канал и максимум два дополнительных.","Публиковать загадки, атмосферу, дилеммы и реальный процесс.","Платная реклама — только с измеримой целью и лимитом."]),
("Риски и метрики",["Следить за темпом, логикой магии и отличием от похожих книг.","Измерять запросы фрагмента, подписки, переходы и конверсию.","Не покупать отзывы, не накручивать охваты."]),
("Публикация",["Сравнить традиционное издательство, самостоятельный релиз и сериализацию.","Проверить права, комиссии, территории, сроки и возврат прав.","Выбирать путь по своим приоритетам и реальным условиям."]),
("Следующий шаг",["Уточнить авторское ядро и логлайн.","Создать сильный фрагмент и проверить его на читателях.","Гарантировать успех выше 90% невозможно; можно снижать риски."])]
for n,(title,bullets) in enumerate(slides):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    box=s.shapes.add_textbox(I(.7),I(.5),I(12),I(1)); p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=P(30); p.font.bold=True
    box=s.shapes.add_textbox(I(1),I(1.8),I(11.3),I(5)); tf=box.text_frame; tf.word_wrap=True
    for j,line in enumerate(bullets):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=line; p.font.size=P(22 if len(bullets)<4 else 18); p.space_after=P(18)
prs.save(out/"PROJECT_PRESENTATION.pptx")

wb=Workbook(); ws=wb.active; ws.title="Конкуренты"
ws.append(["Название","Автор","Год","Рынок","Жанр","Логлайн","Обложка","Аннотация","Начало","Похвалы","Жалобы","Публичные данные","Рейтинг и дата","Канал","Отличие","Урок","Источник URL","Дата проверки"])
for _ in range(25): ws.append([""]*18)
for name,headers,rows in [
("Бюджет",["Статья","План","Факт","Когда нужна","Условие"],[["Редактура","","","После черновика","Сравнить предложения"],["Корректура","","","Перед релизом","Финальный текст"],["Обложка","","","Перед релизом","Проверить позиционирование"],["Реклама","","","После теста страницы","Задать лимит и метрику"],["Резерв","","","Всегда","Не тратить на непроверенное"]]),
("Эксперименты",["Дата","Гипотеза","Материал","Аудитория","Расход","Охват","Ответы","Подписки","Переходы","Покупки","Вывод"],[[""]*11 for _ in range(30)]),
("30 дней",["День","Действие","Результат","Выполнено"],[[i,x,"",""] for i,x in enumerate(["3 логлайна","Определить финальное чувство","Список 15–25 книг","Изучить пять книг","Описать сегменты","Проверить логлайны","Выбрать канал","Опубликовать концепцию","Собрать обратную связь","Отредактировать фрагмент","Проверить визуалы","Собрать пресс-пакет","Изучить правила сообществ","Оценить отклик","Записать расходы","Решить следующий цикл"],1)])]:
    sh=wb.create_sheet(name); sh.append(headers)
    for row in rows: sh.append(row)
for sh in wb.worksheets:
    sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="243B53"); cell.alignment=Alignment(wrap_text=True)
    for col in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(col)].width=20
wb.save(out/"PROJECT_TRACKERS.xlsx")

# Word version of the competitor research CSV template
template_doc=Document()
p=template_doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Шаблон анализа конкурентов")
template_doc.add_paragraph("Используйте эту форму для 15–25 сопоставимых книг. Заполняйте только проверяемые сведения, фиксируйте дату и источник; неизвестные продажи не угадывайте.")
fields=[
("Библиографические данные","Название, автор, год публикации, язык и целевой рынок."),
("Жанр и обещание","Жанр, логлайн/центральная идея, обещание обложки и аннотации."),
("Начало книги","Как устроена первая зацепка; в какой момент появляется центральный вопрос."),
("Реакция читателей","Повторяющиеся похвалы и жалобы из доступных отзывов. Не считайте отдельный отзыв репрезентативным."),
("Публичные показатели","Только опубликованные данные продаж, рейтинги с датой, награды и доступные сведения о тиражах."),
("Обнаружение книги","Где читатели могли узнать о книге: рекомендации, видео, блог, магазин, библиотека или другое."),
("Отличие и урок","Что делает книгу узнаваемой; что можно изучить, не копируя чужой текст или оформление."),
("Источники","URL, дата проверки, оговорки и уровень уверенности.")
]
for title,desc in fields:
    template_doc.add_heading(title,level=2); template_doc.add_paragraph(desc)
template_doc.add_heading("Карточка одной книги",level=1)
for label in ["Название и автор","Год / язык / рынок","Жанр","Логлайн","Обложка и аннотация","Зацепка первых страниц","Похвалы читателей","Жалобы читателей","Публичные данные и дата","Канал обнаружения","Отличие","Урок для собственного проекта","Источник URL и дата проверки"]:
    p=template_doc.add_paragraph(); r=p.add_run(label+": "); r.bold=True; p.add_run("____________________________________________________________")
template_doc.save(out/"10_COMPETITOR_RESEARCH_TEMPLATE.docx")

files=sorted(x.name for x in out.iterdir() if x.is_file() and not x.name.startswith("."))
(out/"README_OUTPUT.md").write_text("# Готовые документы\n\n" + "\n".join("- ["+x+"]("+x+")" for x in files)+"\n",encoding="utf-8")
print("Generated",len(files),"files")
